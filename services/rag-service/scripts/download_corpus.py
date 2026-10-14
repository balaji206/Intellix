"""Download the Phase 2 evaluation corpus into evaluation/datasets/corpus/.

Place this file in services/rag-service/scripts/ and run from services/rag-service:

    pip install beautifulsoup4
    python -m scripts.download_corpus            # skips files that already exist
    python -m scripts.download_corpus --force    # re-download everything

What it does
- Downloads 2 PDFs and 2 RFC text files directly.
- Downloads 3 GitLab Handbook pages, converts them to clean text, and saves them as
  Markdown (.md) and Word (.docx) so every parser gets exercised.
- Runs each file through the project's own parser and prints how many sections and
  characters it extracted, so you know right away whether a file is usable.

See evaluation/datasets/SOURCES.md for sources and licence notes.
"""
import sys
from pathlib import Path

import httpx

ROOT = Path(__file__).resolve().parents[3]
OUT = ROOT / "evaluation" / "datasets" / "corpus"

HEADERS = {
    "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) Intellix-eval-corpus-downloader",
    "Accept": "*/*",
}

# (filename, [candidate URLs tried in order])
DOWNLOADS = [
    (
        "NIST.AI.100-1.pdf",
        ["https://nvlpubs.nist.gov/nistpubs/ai/NIST.AI.100-1.pdf"],
    ),
    (
        "DPDP-Act-2023.pdf",
        [
            "https://www.indiacode.nic.in/bitstream/123456789/22037/2/a2023-22.pdf",
            "https://www.meity.gov.in/writereaddata/files/Digital%20Personal%20Data%20Protection%20Act%202023.pdf",
        ],
    ),
    (
        "rfc7519-jwt.txt",
        [
            "https://www.ietf.org/rfc/rfc7519.txt",
            "https://www.rfc-editor.org/rfc/rfc7519.txt",
        ],
    ),
    (
        "rfc8725-jwt-bcp.txt",
        [
            "https://www.rfc-editor.org/rfc/rfc8725.txt",
            "https://www.ietf.org/rfc/rfc8725.txt",
        ],
    ),
]

# (filename, URL, output format)
PAGES = [
    (
        "gitlab-time-off-types.md",
        "https://handbook.gitlab.com/handbook/people-group/time-off-and-absence/time-off-types/",
        "md",
    ),
    (
        "gitlab-leave-types.docx",
        "https://handbook.gitlab.com/handbook/people-group/time-off-and-absence/leave-types/",
        "docx",
    ),
    (
        "gitlab-time-away-philosophy.md",
        "https://handbook.gitlab.com/handbook/people-group/time-off-and-absence/time-away-philosophy/",
        "md",
    ),
]


# ---------------------------------------------------------------- HTML -> blocks
def html_to_blocks(html: str) -> list[tuple[str, str]]:
    """Return a list of (kind, text). kind is h1..h4, li, tr or p."""
    from bs4 import BeautifulSoup

    soup = BeautifulSoup(html, "html.parser")
    for tag in soup(["script", "style", "nav", "footer", "header", "aside", "form", "noscript", "svg"]):
        tag.decompose()

    root = soup.find("main") or soup.find("article") or soup.body or soup
    blocks: list[tuple[str, str]] = []

    for el in root.find_all(["h1", "h2", "h3", "h4", "p", "li", "tr", "pre"]):
        # avoid duplicating text that a parent block already covers
        if el.name in ("p", "pre") and el.find_parent(["li", "td", "th"]):
            continue
        if el.name == "li" and el.find_parent(["td", "th"]):
            continue

        if el.name == "tr":
            cells = [" ".join(c.get_text(" ", strip=True).split()) for c in el.find_all(["th", "td"])]
            text = " | ".join(c for c in cells if c)
        elif el.name == "li":
            clone = BeautifulSoup(str(el), "html.parser")
            for sub in clone.find_all(["ul", "ol"]):
                sub.decompose()
            text = " ".join(clone.get_text(" ", strip=True).split())
        else:
            text = " ".join(el.get_text(" ", strip=True).split())

        if not text:
            continue
        if blocks and blocks[-1] == (el.name, text):
            continue
        blocks.append((el.name, text))
    return blocks


def write_markdown(blocks: list[tuple[str, str]], path: Path) -> None:
    lines = []
    for kind, text in blocks:
        if kind in ("h1", "h2", "h3", "h4"):
            lines.append("#" * int(kind[1]) + " " + text)
        elif kind == "li":
            lines.append("- " + text)
        elif kind == "tr":
            lines.append("| " + text + " |")
        else:
            lines.append(text)
    path.write_text("\n\n".join(lines) + "\n", encoding="utf-8")


def write_docx(blocks: list[tuple[str, str]], path: Path) -> None:
    from docx import Document

    doc = Document()
    for kind, text in blocks:
        if kind in ("h1", "h2", "h3", "h4"):
            doc.add_heading(text, level=int(kind[1]))
        elif kind == "li":
            doc.add_paragraph(text, style="List Bullet")
        else:
            doc.add_paragraph(text)
    doc.save(str(path))


# ---------------------------------------------------------------- downloading
def fetch(client: httpx.Client, urls: list[str]) -> tuple[bytes, str]:
    last_error = "no URL tried"
    for url in urls:
        try:
            r = client.get(url)
            if r.status_code == 200 and r.content:
                return r.content, url
            last_error = f"{url} -> HTTP {r.status_code}"
        except Exception as e:  # network, DNS, SSL ...
            last_error = f"{url} -> {type(e).__name__}: {e}"
    raise RuntimeError(last_error)


def check_parse(path: Path) -> str:
    try:
        from app.services.parser import parse_file

        sections = parse_file(path)
        chars = sum(len(s.text) for s in sections)
        if not sections:
            return "WARNING: parser extracted no text"
        return f"parsed OK: {len(sections)} section(s), {chars:,} chars"
    except Exception as e:
        return f"parse check failed: {type(e).__name__}: {e}"


def main() -> int:
    force = "--force" in sys.argv
    OUT.mkdir(parents=True, exist_ok=True)
    results: list[tuple[str, str]] = []

    with httpx.Client(headers=HEADERS, follow_redirects=True, timeout=90) as client:
        for name, urls in DOWNLOADS:
            path = OUT / name
            if path.exists() and not force:
                results.append((name, "already exists, " + check_parse(path)))
                continue
            try:
                data, used = fetch(client, urls)
                if name.endswith(".pdf") and not data.startswith(b"%PDF"):
                    raise RuntimeError(f"{used} did not return a PDF (blocked or redirected?)")
                path.write_bytes(data)
                results.append((name, check_parse(path)))
            except Exception as e:
                results.append((name, f"FAILED ({e}). Download it manually and save as {path}"))

        for name, url, fmt in PAGES:
            path = OUT / name
            if path.exists() and not force:
                results.append((name, "already exists, " + check_parse(path)))
                continue
            try:
                data, used = fetch(client, [url])
                blocks = html_to_blocks(data.decode("utf-8", errors="ignore"))
                if len(blocks) < 5:
                    raise RuntimeError("page produced almost no text (layout may have changed)")
                (write_markdown if fmt == "md" else write_docx)(blocks, path)
                results.append((name, check_parse(path)))
            except Exception as e:
                results.append((name, f"FAILED ({e}). Open {url}, copy the text, and save it as {path}"))

    print(f"\nCorpus folder: {OUT}\n")
    width = max(len(n) for n, _ in results)
    for name, status in results:
        print(f"{name.ljust(width)}  {status}")

    failed = [n for n, s in results if s.startswith("FAILED") or s.startswith("WARNING")]
    print(f"\n{len(results) - len(failed)}/{len(results)} files ready.")
    return 1 if failed else 0


if __name__ == "__main__":
    raise SystemExit(main())
